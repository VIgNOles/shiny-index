import copy
import os
import shutil
import tempfile
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from scripts.import_sheet_export import review
from src.indexer import accept, empty, save_master


class SheetExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        self.master = root / 'master.xlsx'
        self.export = root / 'sheet-export.xlsx'
        card = {
            'card_title': '【試験】', 'idol_id': 'idol_test', 'idol_name': '試験',
            'card_kind': 'P', 'rarity': 'SR', 'unit_id': None, 'unit_name': None,
            'first_implemented_on': '2020-01-01', 'acquisition_category': 'unknown',
            'series_ids': [], 'series_status': 'unknown', 'collab_work': None,
            'wiki_url': 'https://wikiwiki.jp/shinycolors/試験',
            'wiki_link_status': 'observed', 'variant_kind': 'base',
            'review_status': 'needs_review', 'source_ref': 'test',
        }
        batch = {'status': 'validated', 'cards': [card], 'run_id': 'test',
                 'observed_at': '2020-01-01', 'scope': 'sample'}
        with patch.dict(os.environ, {'XLSX_BACKEND': 'stdlib'}):
            save_master(accept(empty(), batch), self.master)
        shutil.copy2(self.master, self.export)

    def edit(self, tab, row, values):
        book = load_workbook(self.export)
        sheet = book[tab]
        for header, value in values.items():
            column = next(c.column for c in sheet[1] if c.value == header)
            sheet.cell(row, column).value = value
        book.save(self.export)

    def test_manual_change_is_reviewed_without_changing_master(self):
        self.edit('追加・手修正', 2, {
            'card_title': '【確認済】', 'reason': '試験',
            'source_ref': 'manual-test', 'updated_at': '2026-10-07T00:00:00+09:00',
        })
        before = self.master.read_bytes()
        result = review(self.master, self.export)
        self.assertEqual(result['before_count'], 1)
        self.assertEqual(result['after_count'], 1)
        self.assertEqual(len(result['new_edit_ids']), 1)
        self.assertEqual(result['edit_changes'][0]['after']['values']['card_title'], '【確認済】')
        self.assertTrue(result['protected_tabs_unchanged'])
        self.assertEqual(self.master.read_bytes(), before)

    def test_manual_duplicate_requires_existing_id(self):
        self.edit('追加・手修正', 3, {
            'card_id': str(uuid.uuid4()), 'card_title': '【試験】',
            'idol_id': 'idol_test', 'idol_name': '試験', 'card_kind': 'P',
            'variant_kind': 'base', 'reason': '試験',
            'source_ref': 'manual-test', 'updated_at': '2026-10-07T00:00:00+09:00',
        })
        with self.assertRaisesRegex(ValueError, 'possible manual duplicate'):
            review(self.master, self.export)
        self.edit('追加・手修正', 3, {'card_title': '【別カード】'})
        result = review(self.master, self.export)
        self.assertEqual(result['after_count'], 2)
        self.assertEqual(len(result['added_cards']), 1)

    def test_protected_tab_edit_is_rejected(self):
        self.edit('_meta', 2, {'record_json': '{"revision":99}'})
        with self.assertRaisesRegex(ValueError, 'protected tab changed'):
            review(self.master, self.export)


if __name__ == '__main__':
    unittest.main()
