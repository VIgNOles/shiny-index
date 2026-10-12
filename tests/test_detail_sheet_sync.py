import copy
from contextlib import redirect_stdout
import io
import json
import os
import shutil
import sys
import unittest
import uuid
from unittest.mock import patch

from openpyxl import load_workbook

from scripts.import_sheet_export import main, review
from src.detail_master import DETAIL_TABS, adopt, alias, empty, from_workbook, table_rows
from src.indexer import load_master, save_master, write
from tests import test_sheet_import as sheet_import
from tests.test_detail_master import candidate, rehash


class DetailSheetSyncTests(unittest.TestCase):
    def setUp(self):
        sheet_import.SheetExportTests.setUp(self)
        cid = load_master(self.master)['registry'][0]['card_id']
        document = candidate()
        document['cards'][0]['card_id'] = cid
        document['card_coverage'][0]['card_id'] = cid
        rehash(document)
        self.detail = adopt(empty(document['base_dataset_version']), document)
        self.rewrite_details(self.master, self.detail)
        shutil.copy2(self.master, self.export)
        self.report = self.master.parent / 'review.json'

    def rewrite_details(self, path, detail):
        book = load_workbook(path)
        try:
            for name in DETAIL_TABS:
                if name in book:
                    book.remove(book[name])
            for name, rows in table_rows(detail).items():
                sheet = book.create_sheet(name)
                for row in rows:
                    sheet.append(row)
            book.save(path)
        finally:
            book.close()

    def edit(self, tab, row, values):
        sheet_import.SheetExportTests.edit(self, tab, row, values)

    def manual_edit(self, **values):
        self.edit('詳細項目', 2, {
            '手修正_名称': '変更名称', '理由': '訂正', '根拠': 'source',
            '更新日時': '2026-10-09T04:00:00Z', **values,
        })

    def run_apply(self):
        with patch.object(sys, 'argv', ['import_sheet_export.py', 'apply',
                                       str(self.master), str(self.export), str(self.report)]), \
                patch.dict(os.environ, {'XLSX_BACKEND': 'stdlib'}), redirect_stdout(io.StringIO()):
            main()

    def test_detail_only_edit_reports_fields_and_resolved_values(self):
        before = self.master.read_bytes()
        self.manual_edit(**{'手修正_SP': 30, '手修正_効果': 'Vocal3倍アピール (Plus)',
                            '手修正_タグJSON': '["plus"]'})
        report = review(self.master, self.export)
        detail_id = self.detail['registry'][0]['detail_id']
        self.assertTrue(report['detail_changed'])
        self.assertTrue(report['detail_protected_unchanged'])
        self.assertEqual(report['detail_before_count'], 1)
        self.assertEqual(report['detail_after_count'], 1)
        self.assertEqual(report['detail_new_edit_ids'], [detail_id])
        self.assertEqual(report['edit_changes'], [])
        change = report['detail_edit_changes'][0]
        self.assertEqual(change['detail_id'], detail_id)
        self.assertEqual(change['card_id'], self.detail['registry'][0]['card_id'])
        self.assertEqual(change['before']['override'], {})
        self.assertEqual(change['after']['override'], {
            'name': '変更名称', 'sp': 30, 'effect_private': 'Vocal3倍アピール (Plus)', 'mechanics': ['plus'],
        })
        self.assertEqual(change['resolved_before']['name'], 'スキル')
        self.assertEqual(change['resolved_after']['name'], '変更名称')
        fields = {field['field']: field for field in change['field_changes']}
        self.assertEqual(fields['override.sp'], {'field': 'override.sp', 'before': None, 'after': 30})
        self.assertEqual(fields['reason']['after'], '訂正')
        self.assertEqual(self.master.read_bytes(), before)

    def test_changed_and_removed_manual_edits_are_explicit(self):
        old = copy.deepcopy(self.detail)
        row = old['registry'][0]
        row.update(override={'name': '既存訂正'}, reason='既存理由', source_ref='prior-source',
                   updated_at='2026-10-08T04:00:00Z')
        self.rewrite_details(self.master, old)
        shutil.copy2(self.master, self.export)
        self.manual_edit()
        report = review(self.master, self.export)
        self.assertEqual(report['detail_changed_edit_ids'], [row['detail_id']])
        self.assertEqual(report['detail_edit_changes'][0]['resolved_before']['name'], '既存訂正')
        self.edit('詳細項目', 2, {'手修正_名称': None, '理由': None, '根拠': None, '更新日時': None})
        report = review(self.master, self.export)
        self.assertEqual(report['detail_removed_edit_ids'], [row['detail_id']])
        self.assertEqual(report['detail_edit_changes'][0]['resolved_after']['name'], 'スキル')
        self.assertEqual(report['detail_edit_changes'][0]['before']['reason'], '既存理由')
        self.assertIsNone(report['detail_edit_changes'][0]['after']['reason'])

    def test_metadata_only_change_is_reported(self):
        self.edit('詳細項目', 2, {'理由': '確認のみ', '根拠': 'manual-check', '更新日時': '2026-10-09T04:00:00Z'})
        report = review(self.master, self.export)
        self.assertTrue(report['detail_changed'])
        change = report['detail_edit_changes'][0]
        self.assertEqual(change['resolved_before'], change['resolved_after'])
        self.assertEqual({field['field'] for field in change['field_changes']}, {'reason', 'source_ref', 'updated_at'})

    def test_unchanged_details_are_reported_without_edits(self):
        report = review(self.master, self.export)
        self.assertFalse(report['detail_changed'])
        self.assertEqual(report['detail_edit_changes'], [])
        self.assertEqual(report['detail_revision'], {'before': 1, 'after': 1})

    def test_coherent_protected_item_mutations_are_rejected(self):
        mutations = {
            'source_and_display': lambda detail: detail['registry'][0]['source'].update(effect_private='Vocal99倍アピール'),
            'aliases': lambda detail: detail['registry'][0].update(aliases=['replacement-alias']),
            'origin': lambda detail: detail['registry'][0].update(origin='manual'),
            'fixed_id': lambda detail: detail['registry'][0].update(detail_id=str(uuid.uuid4())),
            'removed_item': lambda detail: detail.update(registry=[]),
        }
        for name, mutate in mutations.items():
            with self.subTest(name=name):
                detail = copy.deepcopy(self.detail)
                mutate(detail)
                self.rewrite_details(self.export, detail)
                with self.assertRaisesRegex(ValueError, 'protected detail item|fixed detail IDs'):
                    review(self.master, self.export)

    def test_new_detail_item_requires_separate_adoption(self):
        detail = copy.deepcopy(self.detail)
        row = copy.deepcopy(detail['registry'][0])
        row['detail_id'] = str(uuid.uuid4())
        row['source']['name'] = '追加スキル'
        row['aliases'] = [alias(row['card_id'], row['source'])]
        detail['registry'].append(row)
        self.rewrite_details(self.export, detail)
        with self.assertRaisesRegex(ValueError, 'fixed detail IDs'):
            review(self.master, self.export)

    def test_coherent_detail_state_mutations_are_rejected(self):
        mutations = {
            'card_source': lambda detail: next(iter(detail['cards'].values()))['source'].update(card_title='別カード'),
            'card_override': lambda detail: next(iter(detail['cards'].values())).update(override={'card_title': '別カード'}),
            'coverage': lambda detail: detail['coverage'][0].update(status='tampered'),
            'base_version': lambda detail: detail.update(base_dataset_version='v1-1111111111111111'),
            'source_runs': lambda detail: detail.update(source_runs=[]),
            'schema': lambda detail: detail.update(detail_schema='2.0'),
        }
        for name, mutate in mutations.items():
            with self.subTest(name=name):
                detail = copy.deepcopy(self.detail)
                mutate(detail)
                self.rewrite_details(self.export, detail)
                with self.assertRaisesRegex(ValueError, 'protected detail state'):
                    review(self.master, self.export)

    def test_redundant_visible_and_extra_cells_are_protected(self):
        mutations = [('詳細カード', 'C2', '表示だけを改変'),
                     ('詳細収録範囲', 'C2', '表示だけを改変'),
                     ('詳細項目', 'G1', '異なる手修正列'),
                     ('詳細項目', 'Q2', '無視される追加値'),
                     ('_detail_meta', 'C2', '無視される追加値')]
        for tab, cell, value in mutations:
            with self.subTest(tab=tab, cell=cell):
                shutil.copy2(self.master, self.export)
                book = load_workbook(self.export)
                book[tab][cell] = value
                book.save(self.export)
                book.close()
                with self.assertRaisesRegex(ValueError, 'protected detail cells'):
                    review(self.master, self.export)

    def test_revision_can_advance_once_with_manual_changes(self):
        self.manual_edit()
        detail = from_workbook(self.export)
        detail['revision'] += 1
        self.rewrite_details(self.export, detail)
        report = review(self.master, self.export)
        self.assertEqual(report['detail_revision'], {'before': 1, 'after': 2})
        for value in [0, 3, True, '2']:
            with self.subTest(revision=value):
                detail['revision'] = value
                self.rewrite_details(self.export, detail)
                with self.assertRaisesRegex(ValueError, 'detail revision'):
                    review(self.master, self.export)

    def test_revision_only_change_is_rejected(self):
        detail = copy.deepcopy(self.detail)
        detail['revision'] += 1
        self.rewrite_details(self.export, detail)
        with self.assertRaisesRegex(ValueError, 'detail revision'):
            review(self.master, self.export)

    def test_formulas_are_rejected(self):
        self.manual_edit(**{'手修正_名称': '=1+1'})
        with self.assertRaisesRegex(ValueError, 'formulas not allowed'):
            review(self.master, self.export)

    def test_missing_detail_tabs_are_rejected(self):
        book = load_workbook(self.export)
        for name in DETAIL_TABS:
            book.remove(book[name])
        book.save(self.export)
        book.close()
        with self.assertRaisesRegex(ValueError, 'Detail tabs removed'):
            review(self.master, self.export)

    def test_added_detail_tabs_require_separate_adoption(self):
        book = load_workbook(self.master)
        for name in DETAIL_TABS:
            book.remove(book[name])
        book.save(self.master)
        book.close()
        with self.assertRaisesRegex(ValueError, 'Detail tabs added'):
            review(self.master, self.export)

    def test_apply_preserves_ids_source_and_existing_manual_edits(self):
        detail = copy.deepcopy(self.detail)
        existing = detail['registry'][0]
        existing.update(override={'name': '既存訂正'}, reason='既存理由', source_ref='prior-source',
                        updated_at='2026-10-08T04:00:00Z')
        second = copy.deepcopy(existing)
        second.update(detail_id=str(uuid.uuid4()), override={})
        second['source']['name'] = '別スキル'
        second['aliases'] = [alias(second['card_id'], second['source'])]
        detail['registry'].append(second)
        self.rewrite_details(self.master, detail)
        shutil.copy2(self.master, self.export)
        self.edit('詳細項目', 3, {'手修正_効果': 'Vocal4倍アピール', '理由': '追加訂正',
                               '根拠': 'new-source', '更新日時': '2026-10-09T04:00:00Z'})
        expected = from_workbook(self.export)
        original_bytes = self.master.read_bytes()
        write(self.report, review(self.master, self.export))
        self.run_apply()
        self.assertEqual(from_workbook(self.master), expected)
        self.assertEqual(from_workbook(self.master)['registry'][0], existing)
        self.assertEqual(load_master(self.master), load_master(self.export))
        backups = list(self.master.parent.glob('master.backup-*.xlsx'))
        self.assertEqual(len(backups), 1)
        self.assertEqual(backups[0].read_bytes(), original_bytes)

    def test_apply_rejects_export_changed_since_review(self):
        self.manual_edit()
        write(self.report, review(self.master, self.export))
        self.manual_edit(**{'手修正_名称': 'レビュー後変更'})
        before = self.master.read_bytes()
        with self.assertRaisesRegex(ValueError, 'changed since review'):
            self.run_apply()
        self.assertEqual(self.master.read_bytes(), before)
        self.assertEqual(list(self.master.parent.glob('master.backup-*.xlsx')), [])

    def test_apply_rejects_master_changed_since_review(self):
        self.manual_edit()
        write(self.report, review(self.master, self.export))
        # A workbook metadata change still invalidates the exact reviewed input.
        book = load_workbook(self.master)
        book.properties.title = 'review後変更'
        book.save(self.master)
        book.close()
        before = self.master.read_bytes()
        with self.assertRaisesRegex(ValueError, 'changed since review'):
            self.run_apply()
        self.assertEqual(self.master.read_bytes(), before)

    def test_apply_without_changes_does_not_rewrite_master(self):
        write(self.report, review(self.master, self.export))
        before = self.master.read_bytes()
        self.run_apply()
        self.assertEqual(self.master.read_bytes(), before)
        self.assertEqual(list(self.master.parent.glob('master.backup-*.xlsx')), [])


if __name__ == '__main__':
    unittest.main()
