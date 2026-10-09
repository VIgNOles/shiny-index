import copy,tempfile,unittest
from pathlib import Path
from openpyxl import Workbook,load_workbook
from src.detail_master import adopt,empty,table_rows,preserve_tabs,from_workbook
from tests.test_detail_master import candidate

class DetailWorkbookTests(unittest.TestCase):
 def test_base_update_preserves_six_detail_tabs(self):
  c=candidate();m=adopt(empty(c['base_dataset_version']),c)
  with tempfile.TemporaryDirectory() as directory:
   source=Path(directory)/'source.xlsx';target=Path(directory)/'target.xlsx'
   wb=Workbook();wb.active.title='基礎タブ'
   for name,rows in table_rows(m).items():
    ws=wb.create_sheet(name)
    for row in rows:ws.append(row)
    ws.freeze_panes='A2';ws.column_dimensions['D'].width=27
   wb.save(source);wb.close()
   wb=Workbook();wb.active.title='基礎タブ';wb.active['A1']='新しい基礎情報';wb.save(target);wb.close()
   preserve_tabs(source,target)
   self.assertEqual(from_workbook(target),m)
   result=load_workbook(target);self.assertEqual(result['基礎タブ']['A1'].value,'新しい基礎情報');self.assertEqual(result['詳細項目'].freeze_panes,'A2');result.close()
 def test_formulas_are_rejected(self):
  c=candidate();m=adopt(empty(c['base_dataset_version']),c)
  with tempfile.TemporaryDirectory() as directory:
   p=Path(directory)/'source.xlsx';wb=Workbook();wb.remove(wb.active)
   for name,rows in table_rows(m).items():
    ws=wb.create_sheet(name)
    for row in rows:ws.append(row)
   wb['詳細項目']['G2']='=1+1';wb.save(p);wb.close()
   with self.assertRaisesRegex(ValueError,'Formula'):from_workbook(p)
if __name__=='__main__':unittest.main()
