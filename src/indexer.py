"""Offline-first card pipeline. Network acquisition never mutates the master."""
import argparse
import copy
import csv
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import uuid
from datetime import date, datetime, timezone
from urllib.parse import urlsplit, urlunsplit, unquote
from urllib.request import Request, urlopen
from urllib.robotparser import RobotFileParser

ROOT = Path(__file__).resolve().parents[1]
FIELDS = ['card_title','idol_id','idol_name','card_kind','rarity','unit_id','unit_name','first_implemented_on','acquisition_category','series_ids','series_status','collab_work','wiki_url','wiki_link_status','variant_kind','review_status']
REQUIRED = ['card_title','idol_id','idol_name','card_kind','variant_kind']
CATEGORIES = ['permanent_gacha','limited_gacha','collection_gacha','collaboration_gacha','event_reward','campaign','mission','gacha_bonus','other_bonus','initial','exchange','other','unknown']
TABS = ['追加・手修正','取得候補','カード確認','辞書','_registry','_source','_evidence','_meta','_runs']

def now(): return datetime.now(timezone.utc).isoformat(timespec='seconds')
def canonical(obj): return json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(',',':'))
def digest(obj): return hashlib.sha256(canonical(obj).encode()).hexdigest()
def read(path): return json.loads(Path(path).read_text(encoding='utf-8'))
def write(path, obj):
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + '.tmp')
    tmp.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')
    os.replace(tmp, path)

def key(card):
    if card.get('source_key'): return card['source_key']
    p = urlsplit(card['wiki_url'])
    return 'wikiwiki:' + unquote(urlunsplit((p.scheme,p.netloc,p.path,'',''))) + ':' + card['variant_kind']

def validate(cards):
    ids=set(); aliases=set()
    for c in cards:
        for f in REQUIRED:
            if not isinstance(c.get(f),str) or not c[f].strip(): raise ValueError(f'missing {f}')
        uuid.UUID(c['card_id']); uuid.UUID(c['family_id'])
        if c['card_id'] in ids: raise ValueError('duplicate ID')
        ids.add(c['card_id'])
        if c['game'] != 'shinycolors_enza' or c['card_kind'] not in ['P','S']: raise ValueError('wrong game/kind')
        if c.get('rarity') not in [None,'N','R','SR','SSR','UR']: raise ValueError('unknown rarity')
        if c['variant_kind'] not in ['base','idol_road_sr','idol_road_ssr','other']: raise ValueError('variant')
        if c.get('acquisition_category') not in CATEGORIES: raise ValueError('category')
        if c.get('series_status') not in ['known','none','unknown'] or not isinstance(c.get('series_ids'),list): raise ValueError('series')
        if c.get('review_status') not in ['wiki_only','official_checked','user_evidence_checked','needs_review']: raise ValueError('review')
        d=c.get('first_implemented_on')
        if d and (date.fromisoformat(d).isoformat()!=d or d>date.today().isoformat()): raise ValueError('invalid/future date')
        url=c.get('wiki_url')
        if url:
            p=urlsplit(url)
            if p.scheme!='https' or p.netloc!='wikiwiki.jp' or not unquote(p.path).startswith('/shinycolors/'): raise ValueError('unsafe wiki URL')
            k=key(c)
            if k in aliases: raise ValueError('duplicate source variant')
            aliases.add(k)
    if any(c['family_id'] not in ids for c in cards): raise ValueError('orphan family')

def empty(): return {'revision':0,'registry':[],'source':[],'overrides':[],'evidence':[],'runs':[],'dictionaries':[], 'candidates':[], 'coverage':{'scope':'sample','complete':False,'unverified':['全期間のWiki掲載分','ゲーム全網羅','ロード派生','公式独立照合']}}

def resolve(m):
    result=[]
    registered=[r['card_id'] for r in m['registry']]
    aliases=[a for r in m['registry'] for a in r['aliases']]
    if len(registered)!=len(set(registered)) or len(aliases)!=len(set(aliases)):raise ValueError('registry collision')
    if any(o['card_id'] not in registered for o in m['overrides']):raise ValueError('orphan override')
    for reg in m['registry']:
        c={f:None for f in FIELDS}; c.update({'game':'shinycolors_enza','series_ids':[],'series_status':'unknown','acquisition_category':'unknown','variant_kind':'base','review_status':'needs_review','wiki_link_status':'unknown'})
        c.update(next((x['values'] for x in m['source'] if x['card_id']==reg['card_id']),{}))
        override=next((x for x in m['overrides'] if x['card_id']==reg['card_id']),{})
        values=override.get('values',{})
        if set(values)-set(FIELDS): raise ValueError('unknown override fields')
        c.update({k:v for k,v in values.items() if v is not None and v!=''})
        clear=override.get('clear_fields',[])
        if set(clear)-set(FIELDS) or set(clear)&set(values): raise ValueError('invalid/conflicting clear')
        for f in clear: c[f]=None
        c.update({k:reg[k] for k in ['card_id','family_id']})
        c['record_updated_at']=override.get('updated_at') or reg['updated_at']
        c['source_refs']=sorted({x['source_ref'] for x in m['evidence'] if x['card_id']==reg['card_id']})
        if override.get('source_ref'): c['source_refs']=sorted(set(c['source_refs']+[override['source_ref']]))
        c['unknown_fields']=[f for f in FIELDS if c[f] is None or c[f]=='unknown']
        result.append(c)
    validate(result)
    return sorted(result,key=lambda c:c['card_id'])

def accept(m, batch, mappings=None):
    m=copy.deepcopy(m); mappings=mappings or {}
    if batch['status']!='validated' or not batch['cards']: raise ValueError('invalid/empty batch')
    previous=next((r for r in m['runs'] if r['run_id']==batch['run_id']),None)
    if previous:
        if previous.get('batch_hash') and previous['batch_hash']!=digest(batch):raise ValueError('run_id reused with different content')
        return m
    incoming=[key(c) for c in batch['cards']]
    if len(incoming)!=len(set(incoming)): raise ValueError('duplicate incoming source')
    if batch.get('scope')=='full':
        old={a for r in m['registry'] for a in r['aliases']}
        if old-set(incoming): raise ValueError('missing existing source keys; review required')
        if old and (abs(len(incoming)-len(old))>=20 or abs(len(incoming)-len(old))/len(old)>=.05): raise ValueError('count anomaly; review required')
    if batch.get('scope')=='initial-full' and m['coverage'].get('scope')!='sample': raise ValueError('initial full import already done; use reviewed full update')
    resolved=resolve(m)
    for candidate in batch['cards']:
        k=key(candidate)
        reg=next((r for r in m['registry'] if k in r['aliases']),None)
        if reg is None and k in mappings:
            reg=next(r for r in m['registry'] if r['card_id']==mappings[k])
            existing=next(c for c in resolved if c['card_id']==reg['card_id'])
            if any(existing[f]!=candidate[f] for f in ['idol_id','card_kind','variant_kind']): raise ValueError('unsafe identity mapping')
            reg['aliases'].append(k)
        if reg is None:
            similar=[c for c in resolved if all(c[f]==candidate[f] for f in ['idol_id','card_kind','variant_kind','card_title'])]
            if similar: raise ValueError('possible manual duplicate; explicit mapping required')
            cid=str(uuid.uuid4())
            reg={'card_id':cid,'family_id':cid,'aliases':[k],'updated_at':now()}
            m['registry'].append(reg)
        values={f:candidate.get(f) for f in FIELDS}
        old=next((x for x in m['source'] if x['card_id']==reg['card_id']),None)
        if old is None or old['values']!=values:
            m['source']=[x for x in m['source'] if x['card_id']!=reg['card_id']]
            m['source'].append({'card_id':reg['card_id'],'values':values,'accepted_run_id':batch['run_id'],'raw_acquisition':candidate.get('raw_acquisition')})
            reg['updated_at']=now()
            m['evidence']=[x for x in m['evidence'] if x['card_id']!=reg['card_id']]
            for f in FIELDS:
                if values[f] is not None:
                    e={'card_id':reg['card_id'],'field_name':f,'value':values[f],'source_ref':candidate['source_ref'],'url':candidate.get('source_page') or candidate['wiki_url'],'observed_at':batch['observed_at'],'verification_status':'wiki_only'}
                    if candidate.get('source_locator'):e.update(locator=candidate['source_locator'],response_hash=candidate['source_hash'])
                    m['evidence'].append(e)
    if batch.get('coverage'):m['coverage']=batch['coverage']
    if batch.get('dictionaries'):m['dictionaries']=batch['dictionaries']
    # Link road variants to the base registry after the complete batch is allocated.
    for reg in m['registry']:
        for alias in reg['aliases']:
            if alias.endswith((':idol_road_sr',':idol_road_ssr')):
                base=alias.rsplit(':',1)[0]+':base'
                parent=next((r for r in m['registry'] if base in r['aliases']),None)
                if parent is None:raise ValueError('orphan road variant')
                reg['family_id']=parent['card_id']
    m['revision']+=1; m['runs'].append({'run_id':batch['run_id'],'accepted_at':now(),'batch_hash':digest(batch)})
    resolve(m)
    return m

def workbook_spec(m):
    cards=resolve(m)
    headers=['card_id']+FIELDS+['clear_fields','reason','source_ref','updated_at']
    edits=[]
    for reg in m['registry']:
        ov=next((x for x in m['overrides'] if x['card_id']==reg['card_id']),{})
        edits.append([reg['card_id']]+[ov.get('values',{}).get(f) for f in FIELDS]+[ov.get('clear_fields',[]),ov.get('reason'),ov.get('source_ref'),ov.get('updated_at')])
    sheets=[{'name':'追加・手修正','rows':[headers]+edits}, {'name':'取得候補','rows':[['run_id','status','details']]+[[x.get('run_id'),x.get('status'),x] for x in m['candidates']]}, {'name':'カード確認','rows':[['card_id']+FIELDS]+[[c['card_id']]+[c.get(f) for f in FIELDS] for c in cards]}]
    for name, data in [('辞書',m['dictionaries']),('_registry',m['registry']),('_source',m['source']),('_evidence',m['evidence']),('_meta',[{'revision':m['revision'],'coverage':m['coverage']}]),('_runs',m['runs'])]:
        sheets.append({'name':name,'rows':[['record_json']]+[[x] for x in data]})
    return {'sheets':sheets}

def load_master(path):
    from openpyxl import load_workbook
    wb=load_workbook(path,data_only=False)
    if any(c.data_type=='f' for ws in wb for row in ws for c in row):raise ValueError('formulas not allowed in canonical master')
    if set(TABS)-set(wb.sheetnames): raise ValueError('missing master tabs')
    def records(name): return [json.loads(row[0]) for row in wb[name].iter_rows(min_row=2,values_only=True) if row[0]]
    meta=records('_meta')[0]
    m={'revision':meta['revision'],'coverage':meta['coverage'],'registry':records('_registry'),'source':records('_source'),'evidence':records('_evidence'),'runs':records('_runs'),'dictionaries':records('辞書'),'candidates':[],'overrides':[]}
    rows=list(wb['追加・手修正'].iter_rows(values_only=True)); headers=list(rows[0]); seen=set()
    for row in rows[1:]:
        x=dict(zip(headers,row)); cid=x.get('card_id')
        if not any(v is not None for v in row): continue
        if not cid or cid in seen: raise ValueError('missing/duplicate edit ID')
        seen.add(cid)
        vals={f:x[f] for f in FIELDS if x.get(f) is not None}
        if 'series_ids' in vals: vals['series_ids']=json.loads(vals['series_ids'])
        clear=json.loads(x.get('clear_fields') or '[]')
        if vals or clear:
            if not x.get('reason') or not x.get('source_ref'): raise ValueError('edit requires reason/source_ref')
            if not any(r['card_id']==cid for r in m['registry']):
                uuid.UUID(cid); m['registry'].append({'card_id':cid,'family_id':cid,'aliases':[],'updated_at':now()})
            if not x.get('updated_at'): raise ValueError('edit requires updated_at (ISO date/time)')
            m['overrides'].append({'card_id':cid,'values':vals,'clear_fields':clear,'reason':x['reason'],'source_ref':x['source_ref'],'updated_at':x['updated_at']})
    resolve(m); return m

def xlsx(spec, path):
    if os.environ.get('XLSX_BACKEND')=='stdlib':
        from src.xlsx_fallback import build
        build(spec,path)
        return
    specpath=Path(str(path)+'.spec.json'); write(specpath,spec)
    subprocess.run([os.environ.get('NODE','node'),str(ROOT/'scripts/workbook.mjs'),str(specpath),str(path)],check=True)
    # Artifact Tool coerces ISO strings to Excel dates even with @ formatting.
    # Repair only cell storage types in the exported OOXML; retain authored layout.
    import zipfile
    import xml.etree.ElementTree as ET
    ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ET.register_namespace('',ns)
    def colnum(label):
        n=0
        for ch in label: n=n*26+ord(ch)-64
        return n-1
    import re
    tmp=Path(str(path)+'.typed')
    with zipfile.ZipFile(path) as z, zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as out:
        for info in z.infolist():
            data=z.read(info.filename)
            match=re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml',info.filename)
            if match:
                rows=spec['sheets'][int(match[1])-1]['rows']; root=ET.fromstring(data)
                for c in root.findall('.//{'+ns+'}c'):
                    co,ro=re.fullmatch(r'([A-Z]+)(\d+)',c.attrib['r']).groups(); ri=int(ro)-1; ci=colnum(co)
                    if ri>=len(rows) or ci>=len(rows[ri]): continue
                    v=rows[ri][ci]
                    if isinstance(v,(dict,list)): v=json.dumps(v,ensure_ascii=False,separators=(',',':'))
                    if isinstance(v,str):
                        for child in list(c): c.remove(child)
                        c.set('t','inlineStr'); isel=ET.SubElement(c,'{'+ns+'}is'); t=ET.SubElement(isel,'{'+ns+'}t'); t.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); t.text=v
                data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            out.writestr(info,data)
    os.replace(tmp,path)

def save_master(m,path):
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    if path.exists(): shutil.copy2(path,path.with_name(path.stem+'.backup-'+datetime.now().strftime('%Y%m%d%H%M%S%f')+'.xlsx'))
    staged=path.with_name(path.stem+'.staged.xlsx'); xlsx(workbook_spec(m),staged)
    check=load_master(staged)
    if resolve(check)!=resolve(m): raise ValueError('master roundtrip mismatch')
    os.replace(staged,path)

def prepare(m, output):
    cards=resolve(m)
    coverage=copy.deepcopy(m['coverage']); dates=[c['first_implemented_on'] for c in cards if c['first_implemented_on']]
    coverage.update({'count':len(cards),'families':len({c['family_id'] for c in cards}),'min_date':min(dates) if dates else None,'max_date':max(dates) if dates else None,'duplicate_keys':0,'by_kind':{k:sum(c['card_kind']==k for c in cards) for k in ['P','S']},'missing':{f:sum(c.get(f) is None for c in cards) for f in FIELDS}})
    from collections import Counter
    coverage['by_rarity']=dict(Counter(c['card_kind']+'-'+str(c['rarity']) for c in cards))
    coverage['by_year']=dict(sorted(Counter((c['first_implemented_on'] or 'unknown')[:4] for c in cards).items()))
    coverage['by_unit']=dict(Counter(c['unit_name'] or 'unknown' for c in cards))
    coverage['by_category']=dict(Counter(c['acquisition_category'] for c in cards))
    coverage['wiki_unavailable']=sum(c['wiki_url'] is None for c in cards)
    sources=copy.deepcopy(m['evidence'])
    for ov in m['overrides']:
        for f in set(ov['values'])|set(ov.get('clear_fields',[])):
            sources=[e for e in sources if not(e['card_id']==ov['card_id'] and e['field_name']==f)]
            sources.append({'card_id':ov['card_id'],'field_name':f,'value':ov['values'].get(f),'source_ref':ov['source_ref'],'url':None,'observed_at':ov['updated_at'],'verification_status':'manual_assertion'})
    sources=sorted(sources,key=lambda e:(e['card_id'],e['field_name'],e['source_ref']))
    content={'cards':cards,'sources':sources,'coverage':coverage,'redirects':{},'dictionaries':m['dictionaries']}
    h=digest(content); version='v1-'+h[:16]
    out=Path(output); target=out/'data'/version
    if target.exists():
        verify_bundle(target)
        for p in (ROOT/'web').iterdir():
            if p.is_file():
                if p.name=='index.html': (out/p.name).write_text(p.read_text(encoding='utf-8').replace('__VERSION__',version),encoding='utf-8')
                else: shutil.copy2(p,out/p.name)
        write(out/'data/latest.json',{'dataset_version':version,'manifest':version+'/manifest.json'})
        return version
    meta={'dataset_version':version,'schema_version':'1.0','published_at':now(),'content_hash':h,'input_hash':digest(m),'count':len(cards)}
    staging=Path(tempfile.mkdtemp(prefix='bundle-',dir=out.parent))
    try:
        d=staging/'data'/version; d.mkdir(parents=True)
        write(d/'cards.json',{'meta':meta,**content})
        for name in ['sources','coverage','redirects']: write(d/(name+'.json'),content[name])
        columns=['dataset_version','schema_version','published_at']+list(cards[0]) if cards else []
        rows=[{**{k:meta[k] for k in columns[:3]},**c} for c in cards]
        def cell(v): return canonical(v) if isinstance(v,(list,dict)) else '' if v is None else str(v)
        def safe(v):
            v=cell(v); return "'"+v if v and (v[0] in "=+-@\t\r'" or v.lstrip().startswith(('=','+','-','@'))) else v
        with (d/'cards.csv').open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.writer(f); w.writerow(columns); w.writerows([[safe(r[k]) for k in columns] for r in rows])
        xlsx({'sheets':[{'name':'Cards','rows':[columns]+[[r[k] for k in columns] for r in rows]},{'name':'Metadata','rows':[['key','value']]+list(meta.items())+[['CSV escaping','Leading apostrophe added for formula-like or apostrophe-prefixed strings; remove exactly one. JSON is authoritative.']]},{'name':'Sources','rows':[['card_id','field_name','source_ref','url','observed_at']]+[[e[k] for k in ['card_id','field_name','source_ref','url','observed_at']] for e in sources]}]},d/'cards.xlsx')
        (d/'cards.xlsx.spec.json').unlink(missing_ok=True)
        for log in d.glob('*.inspect.ndjson'): log.unlink()
        manifest={**meta,'files':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in d.iterdir() if p.is_file()}}
        write(d/'manifest.json',manifest)
        verify_bundle(d)
        if out.exists(): shutil.copytree(out,staging,dirs_exist_ok=True)
        for p in (ROOT/'web').iterdir():
            if p.is_file(): shutil.copy2(p,staging/p.name)
        html=(staging/'index.html').read_text(encoding='utf-8').replace('__VERSION__',version)
        (staging/'index.html').write_text(html,encoding='utf-8')
        write(staging/'data/latest.json',{'dataset_version':version,'manifest':version+'/manifest.json'})
        out.mkdir(parents=True,exist_ok=True)
        # Local prepared artifact only. Deployment swaps the entire Pages artifact.
        shutil.copytree(staging,out,dirs_exist_ok=True)
    finally:
        assert staging.resolve().is_relative_to(out.parent.resolve())
        shutil.rmtree(staging)
    return version

def verify_bundle(d):
    from openpyxl import load_workbook
    d=Path(d); doc=read(d/'cards.json'); cards=doc['cards']; validate(cards)
    if digest({k:doc[k] for k in ['cards','sources','coverage','redirects','dictionaries']})!=doc['meta']['content_hash']: raise ValueError('content hash')
    manifest=read(d/'manifest.json')
    for name,h in manifest['files'].items():
        if hashlib.sha256((d/name).read_bytes()).hexdigest()!=h: raise ValueError('file hash '+name)
    def val(v): return canonical(v) if isinstance(v,(list,dict)) else '' if v is None else str(v)
    with (d/'cards.csv').open(encoding='utf-8-sig',newline='') as f: csvrows=list(csv.DictReader(f))
    wb=load_workbook(d/'cards.xlsx'); rows=list(wb['Cards'].iter_rows(values_only=True)); excel=[dict(zip(rows[0],r)) for r in rows[1:]]
    if len(csvrows)!=len(cards) or len(excel)!=len(cards): raise ValueError('export count')
    for c,r,e in zip(cards,csvrows,excel):
        for k,v in {**{x:doc['meta'][x] for x in ['dataset_version','schema_version','published_at']},**c}.items():
            cv=r[k][1:] if r[k].startswith("'") else r[k]
            if cv!=val(v) or val(e[k])!=val(v): raise ValueError(f'export mismatch {k}: csv={cv!r} xlsx={e[k]!r} expected={val(v)!r}')
    if any(cell.data_type=='f' for ws in wb for row in ws for cell in row): raise ValueError('formula in output')
    allowed={'cards.json','cards.csv','cards.xlsx','sources.json','coverage.json','redirects.json','manifest.json'}
    if {p.name for p in d.iterdir()}!=allowed: raise ValueError('unexpected public file')
    return {'count':len(cards),'version':doc['meta']['dataset_version'],'verified':True}

def rollback(output,version):
    output=Path(output)
    if not version.startswith('v1-') or len(version)!=19 or any(c not in '0123456789abcdef' for c in version[3:]): raise ValueError('invalid version')
    verify_bundle(output/'data'/version)
    html=(ROOT/'web/index.html').read_text(encoding='utf-8').replace('__VERSION__',version)
    tmp=output/'index.html.next';tmp.write_text(html,encoding='utf-8');os.replace(tmp,output/'index.html')
    write(output/'data/latest.json',{'dataset_version':version,'manifest':version+'/manifest.json'})

def collect(url, directory):
    p=urlsplit(url)
    if p.scheme!='https' or p.netloc!='wikiwiki.jp' or p.query or not unquote(p.path).startswith('/shinycolors/'): raise ValueError('URL outside allowlist')
    directory=Path(directory); directory.mkdir(parents=True,exist_ok=False)
    report={'url':url,'fetched_at':now(),'status':'failed'}
    try:
        robots=urlopen('https://wikiwiki.jp/robots.txt',timeout=25).read().decode()
        rp=RobotFileParser(); rp.parse(robots.splitlines())
        if not rp.can_fetch('ShinycolorsCardIndex',url): raise ValueError('robots denies')
        # Percent-encode only Unicode; preserve observed href spelling.
        from urllib.parse import quote
        req=Request(quote(url,safe=':/%'),headers={'User-Agent':'ShinycolorsCardIndex/0.1 (limited research)'})
        with urlopen(req,timeout=25) as r:
            raw=r.read(); report.update({'http_status':r.status,'content_type':r.headers.get('Content-Type')})
        if b'<html' not in raw.lower() or len(raw)<1000: raise ValueError('unexpected response')
        (directory/'response.html').write_bytes(raw)
        report.update({'status':'fetched','sha256':hashlib.sha256(raw).hexdigest()})
    finally: write(directory/'fetch.json',report)

def main():
    p=argparse.ArgumentParser(); sub=p.add_subparsers(dest='cmd',required=True)
    a=sub.add_parser('collect'); a.add_argument('url'); a.add_argument('directory')
    a=sub.add_parser('accept'); a.add_argument('batch'); a.add_argument('master'); a.add_argument('--mapping')
    a=sub.add_parser('prepare'); a.add_argument('master'); a.add_argument('output')
    a=sub.add_parser('verify'); a.add_argument('directory')
    a=sub.add_parser('new-id')
    a=sub.add_parser('rollback'); a.add_argument('output'); a.add_argument('version')
    args=p.parse_args()
    if args.cmd=='collect': collect(args.url,args.directory)
    elif args.cmd=='new-id': print(uuid.uuid4())
    elif args.cmd=='verify': print(canonical(verify_bundle(args.directory)))
    elif args.cmd=='rollback': rollback(args.output,args.version)
    elif args.cmd=='accept': save_master(accept(load_master(args.master) if Path(args.master).exists() else empty(),read(args.batch),read(args.mapping) if args.mapping else None),args.master)
    elif args.cmd=='prepare': print(prepare(load_master(args.master),args.output))

if __name__=='__main__': main()
