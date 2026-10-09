"""Canonical detail registry: stable IDs, source values and manual overrides kept apart."""
import copy
from collections import Counter
from datetime import datetime
import json
import re
import uuid
from unicodedata import normalize
from src.indexer import digest

DETAIL_TABS=['詳細カード','詳細項目','詳細属性','詳細出典','詳細収録範囲','_detail_meta']
OVERRIDE_FIELDS={'name','sp','effect_private','mechanics'}
KINDS={'P':{'panel_live','panel_passive','cap_increase','unique_ability','mb_live','memory_appeal'},
       'S':{'panel_live','panel_passive','cap_increase','unique_ability','quick_skill','possessed_live','support_skill'}}
COLLECTIONS=['panel_nodes','mb_live','memory_appeals','possessed_live','support_skills']


def empty(base_version):
    return {'detail_schema':'1.0','base_dataset_version':base_version,'revision':0,
            'cards':{},'registry':[],'coverage':[],'source_runs':[]}


def alias(card_id,node):
    identity={key:node.get(key) for key in ('kind','sp','unlock_star','unlock_event','mb_stage','level','acquired_at_level')}
    identity['name']=re.sub(r'\s+','',normalize('NFKC',node['name']))
    return card_id+':'+digest(identity)


def resolved_item(row):
    source=row['source'];manual=row.get('override',{})
    item={**copy.deepcopy(source),**manual}
    if 'effect_private' in manual:
        if item['kind'] in {'panel_live','mb_live','possessed_live'} and 'mechanics' not in manual:
            from src.card_details import mechanics
            item['mechanics']=mechanics(item['effect_private'])
        if item['kind']=='cap_increase':
            attribute=r'(?:Vocal|Dance|Visual|メンタル)'
            match=re.fullmatch(rf'({attribute}(?:\s*&\s*{attribute})*)\s*上限\+(\d+)',item['effect_private'])
            if not match:raise ValueError('Manual cap effect must specify targets and increase')
            item.update(cap_targets=[part.strip() for part in match[1].split('&')],cap_delta=int(match[2]))
    return item


def validate(master):
    ids=set();aliases=set()
    coverage_ids=[row['card_id'] for row in master['coverage']]
    if len(coverage_ids)!=len(set(coverage_ids)):raise ValueError('Duplicate coverage card')
    for card_id,card in master['cards'].items():
        uuid.UUID(card_id)
        if card['source']['card_kind'] not in KINDS:raise ValueError('Unknown P/S kind')
    for row in master['registry']:
        uuid.UUID(row['detail_id'])
        if row['detail_id'] in ids or row['card_id'] not in master['cards']:raise ValueError('Duplicate/orphan detail ID')
        ids.add(row['detail_id'])
        if set(row['aliases']) & aliases:raise ValueError('Detail alias collision')
        aliases.update(row['aliases'])
        source=row['source'];manual=row.get('override',{})
        if set(manual)-OVERRIDE_FIELDS:raise ValueError('Immutable or unsupported override field')
        if manual and (not row.get('reason') or not row.get('source_ref') or not row.get('updated_at')):
            raise ValueError('Manual override requires reason, source and date')
        if source['kind'] not in KINDS[master['cards'][row['card_id']]['source']['card_kind']]:raise ValueError('P/S detail mismatch')
        resolved=resolved_item(row)
        if not resolved.get('name'):raise ValueError('Missing detail name')
        if resolved.get('sp') is not None and (type(resolved['sp']) is not int or resolved['sp']<0):raise ValueError('Invalid SP')
        if set(resolved.get('mechanics',[]))-{'link','plus','change','grow','refrain'}:raise ValueError('Invalid mechanics')
        if row.get('updated_at'):
            datetime.fromisoformat(row['updated_at'].replace('Z','+00:00'))
    return master


def adopt(master,candidate,card_ids=None,base_cards=None):
    if candidate['base_dataset_version']!=master['base_dataset_version']:raise ValueError('Base version mismatch')
    if candidate.get('content_hash')!=digest({k:v for k,v in candidate.items() if k!='content_hash'}):raise ValueError('Candidate hash mismatch')
    result=copy.deepcopy(master)
    selected=set(card_ids) if card_ids is not None else {card['card_id'] for card in candidate['cards']}
    if selected-{card['card_id'] for card in candidate['cards']}:raise ValueError('Selected card lacks valid candidate')
    existing={a:row for row in result['registry'] for a in row['aliases']}
    bases={card['card_id']:card for card in (base_cards or [])}
    for card in candidate['cards']:
        cid=card['card_id']
        if cid not in selected:continue
        source={k:v for k,v in card.items() if k not in COLLECTIONS}
        if bases:
            base=bases[cid]
            if base['card_kind']!=card['card_kind']:raise ValueError('Base/candidate kind mismatch')
            source.update(idol_name=base['idol_name'])
        previous=result['cards'].get(cid,{'override':{}})
        if previous.get('source') and previous['source']['card_kind']!=source['card_kind']:raise ValueError('Card kind changed')
        result['cards'][cid]={**previous,'source':source}
        incoming=set();seen=set()
        identity_counts=Counter(alias(cid,node) for collection in COLLECTIONS for node in card.get(collection,[]))
        for collection in COLLECTIONS:
            for node in card.get(collection,[]):
                key=alias(cid,node)
                if identity_counts[key]>1:
                    key+=':cell:'+digest(node['source_positions'][0])
                if key in seen:raise ValueError('Ambiguous duplicate detail identity: '+cid)
                seen.add(key)
                row=existing.get(key)
                if row is None:
                    row={'detail_id':str(uuid.uuid4()),'card_id':cid,'aliases':[key],'source':node,'override':{},'origin':'wiki'}
                    result['registry'].append(row);existing[key]=row
                else:row['source']=copy.deepcopy(node)
                incoming.add(row['detail_id'])
        previous_ids={row['detail_id'] for row in master['registry'] if row['card_id']==cid and row.get('origin')!='manual'}
        if previous_ids-incoming:raise ValueError('Previously registered detail disappeared; review identity mapping')
    result['coverage']=copy.deepcopy(candidate['card_coverage'])
    if candidate['content_hash'] not in result['source_runs']:result['source_runs'].append(candidate['content_hash'])
    result['revision']=master['revision']+1 if result!=master else master['revision']
    return validate(result)


def add_manual(master,card_id,node,reason,source_ref,updated_at):
    result=copy.deepcopy(master)
    if card_id not in result['cards']:raise ValueError('Register the card first')
    key=alias(card_id,node)
    if any(key in row['aliases'] for row in result['registry']):raise ValueError('Manual duplicate; edit existing detail ID')
    result['registry'].append({'detail_id':str(uuid.uuid4()),'card_id':card_id,'aliases':[key],
                             'source':copy.deepcopy(node),'override':{},'origin':'manual',
                             'reason':reason,'source_ref':source_ref,'updated_at':updated_at})
    result['revision']+=1
    return validate(result)


def resolve(master):
    validate(master)
    cards={cid:copy.deepcopy(card['source']) for cid,card in master['cards'].items()}
    for cid,card in master['cards'].items():
        cards[cid].update(card.get('override',{}));cards[cid]['items']=[]
    for row in master['registry']:
        item={**resolved_item(row),'detail_id':row['detail_id']}
        if row.get('override'):item['manual_source_ref']=row['source_ref']
        cards[row['card_id']]['items'].append(item)
    return sorted(cards.values(),key=lambda card:card['card_id'])


def table_rows(master):
    validate(master)
    j=lambda value:json.dumps(value,ensure_ascii=False,separators=(',',':'),sort_keys=True)
    cards=[['card_id','P/S','カード名','アイドル名','アイデア','ひらめき','楽曲熟練度','最大Lv','Vo','Da','Vi','Me','source_json']]
    for cid,record in sorted(master['cards'].items()):
        s=record['source'];traits=s.get('traits',{});stats=s.get('max_status',{})
        cards.append([cid,s['card_kind'],s['card_title'],s.get('idol_name',''),traits.get('idea'),traits.get('inspiration'),
                      j(traits.get('music_proficiencies',[])),stats.get('level'),stats.get('vocal'),stats.get('dance'),stats.get('visual'),stats.get('mental'),j(record)])
    items=[['detail_id','card_id','種類','取得名称','取得SP','取得効果','手修正_名称','手修正_SP','手修正_効果','手修正_タグJSON','理由','根拠','更新日時','source_json','aliases_json','origin']]
    attrs=[['detail_id','card_id','属性','順序','value_json']]
    sources=[['detail_id','card_id','Wiki URL','取得日時','応答SHA256','位置JSON']]
    for row in master['registry']:
        s=row['source'];o=row.get('override',{});card=master['cards'][row['card_id']]['source']
        items.append([row['detail_id'],row['card_id'],s['kind'],s['name'],s.get('sp'),s.get('effect_private'),
                      o.get('name'),o.get('sp'),o.get('effect_private'),j(o['mechanics']) if 'mechanics' in o else None,
                      row.get('reason'),row.get('source_ref'),row.get('updated_at'),j(s),j(row['aliases']),row.get('origin','wiki')])
        for n,step in enumerate(s.get('progression',[])):attrs.append([row['detail_id'],row['card_id'],'support_progression',n,j(step)])
        sources.append([row['detail_id'],row['card_id'],card['wiki_url'],card['fetched_at'],card['source_sha256'],j(s.get('source_positions',[]))])
    coverage=[['card_id','P/S','入力・変換状態','Wiki URL','取得日時','record_json']]
    for row in master['coverage']:
        coverage.append([row['card_id'],row['card_kind'],row['status'],row.get('wiki_url'),row.get('fetched_at'),j(row)])
    meta=[['key','record_json']]+[[key,j(value)] for key,value in master.items() if key not in ('cards','registry','coverage')]
    return dict(zip(DETAIL_TABS,[cards,items,attrs,sources,coverage,meta]))


def from_rows(tables):
    result={key:json.loads(value) for key,value in tables['_detail_meta'][1:] if key}
    result.update(cards={},registry=[],coverage=[])
    for row in tables['詳細カード'][1:]:
        if row[0]:
            if row[0] in result['cards']:raise ValueError('Duplicate detail card')
            result['cards'][row[0]]=json.loads(row[12])
    for row in tables['詳細項目'][1:]:
        row=list(row)+[None]*16
        if not row[0]:continue
        source=json.loads(row[13])
        if row[2]!=source['kind'] or row[3]!=source['name'] or row[4]!=source.get('sp') or row[5]!=source.get('effect_private'):
            raise ValueError('Do not edit fetched columns; use manual columns')
        override={field:row[col] for field,col in [('name',6),('sp',7),('effect_private',8)] if row[col] is not None and row[col]!=''}
        if 'sp' in override:
            value=override['sp']
            if isinstance(value,str) and value.isdecimal():override['sp']=int(value)
            elif type(value) is float and value.is_integer():override['sp']=int(value)
        if row[9]:override['mechanics']=json.loads(row[9])
        result['registry'].append({'detail_id':row[0],'card_id':row[1],'source':source,'aliases':json.loads(row[14]),
                                  'origin':row[15] or 'wiki','override':override,**{key:row[col] for key,col in [('reason',10),('source_ref',11),('updated_at',12)] if row[col]}})
    result['coverage']=[json.loads(row[5]) for row in tables['詳細収録範囲'][1:] if row[0]]
    validate(result)
    expected=table_rows(result)
    for name in ['詳細属性','詳細出典']:
        if tables[name]!=expected[name]:raise ValueError('Derived detail tab changed: '+name)
    return result


def from_workbook(path):
    from openpyxl import load_workbook
    wb=load_workbook(path,data_only=False,read_only=True)
    try:
        tables={}
        for name in DETAIL_TABS:
            rows=[]
            header=list(next(wb[name].iter_rows(min_row=1,max_row=1,values_only=True)))
            while header and header[-1] is None:header.pop()
            width=len(header)
            for cells in wb[name].iter_rows():
                if any(cell.data_type=='f' for cell in cells):raise ValueError('Formula in detail master')
                row=[cell.value for cell in cells]
                if any(value is not None for value in row):rows.append(list(row[:width])+[None]*max(0,width-len(row)))
            tables[name]=rows
        return from_rows(tables)
    finally:wb.close()


def preserve_tabs(source_path,target_path):
    """Base-master updates retain all validated detail values and sheet formatting."""
    from openpyxl import load_workbook
    source=load_workbook(source_path);target=load_workbook(target_path)
    try:
        names=set(source.sheetnames)&set(DETAIL_TABS)
        if not names:return
        if names!=set(DETAIL_TABS):raise ValueError('Incomplete detail tabs; restore before base update')
        original=from_workbook(source_path)
        for name in DETAIL_TABS:
            old=source[name];new=target.create_sheet(name)
            for row in old:
                for cell in row:
                    dest=new.cell(cell.row,cell.column,cell.value)
                    if cell.has_style:dest._style=copy.copy(cell._style)
                    if cell.number_format:dest.number_format=cell.number_format
                    if cell.hyperlink:dest._hyperlink=copy.copy(cell.hyperlink)
                    if cell.comment:dest.comment=copy.copy(cell.comment)
            for key,value in old.column_dimensions.items():new.column_dimensions[key]=copy.copy(value)
            for key,value in old.row_dimensions.items():new.row_dimensions[key]=copy.copy(value)
            new.freeze_panes=old.freeze_panes;new.auto_filter=copy.copy(old.auto_filter)
            new.sheet_format=copy.copy(old.sheet_format);new.sheet_properties=copy.copy(old.sheet_properties)
            new.sheet_view.showGridLines=old.sheet_view.showGridLines
            new.data_validations=copy.deepcopy(old.data_validations)
            for merged in old.merged_cells.ranges:new.merge_cells(str(merged))
        target.save(target_path)
        if from_workbook(target_path)!=original:raise ValueError('Detail tabs changed during base update')
    finally:source.close();target.close()
