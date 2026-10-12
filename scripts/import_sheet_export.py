"""Review and apply a native Google Sheet's XLSX export without changing acquisition data."""
import argparse
import copy
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.indexer import TABS, load_master, now, read, resolve, save_master, write


def file_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def workbook_rows(path, tab_names):
    """Read protected cells with one workbook open, including tab validation."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        from src.detail_master import DETAIL_TABS
        names = set(workbook.sheetnames)
        if not set(TABS) <= names or names - set(TABS) not in (set(), set(DETAIL_TABS)):
            raise ValueError('unexpected master tabs')
        tables = {}
        for name in tab_names:
            rows = []
            for row in workbook[name].iter_rows(values_only=True):
                trimmed = list(row)
                while trimmed and trimmed[-1] is None:
                    trimmed.pop()
                if trimmed:
                    rows.append(trimmed)
            tables[name] = rows
        return tables
    finally:
        workbook.close()


def review_details(master_path, export_path):
    """Only the existing detail items' manual columns may change in this route."""
    from src.detail_master import DETAIL_TABS, OVERRIDE_FIELDS, from_workbook, resolved_item

    def has_details(path):
        book = load_workbook(path, read_only=True)
        try:
            return set(DETAIL_TABS) <= set(book.sheetnames)
        finally:
            book.close()

    had_details, has_incoming = has_details(master_path), has_details(export_path)
    if had_details and not has_incoming:
        raise ValueError('Detail tabs removed from export')
    if has_incoming and not had_details:
        raise ValueError('Detail tabs added; adopt details separately before manual import')
    if not had_details:
        return {
            'detail_changed': False,
            'detail_protected_unchanged': True,
            'detail_before_count': 0,
            'detail_after_count': 0,
            'detail_revision': None,
            'detail_changed_edit_ids': [],
            'detail_removed_edit_ids': [],
            'detail_new_edit_ids': [],
            'detail_edit_changes': [],
        }

    original, incoming = from_workbook(master_path), from_workbook(export_path)
    for field in sorted(original.keys() | incoming.keys()):
        if field not in {'registry', 'revision'} and original.get(field) != incoming.get(field):
            raise ValueError(f'protected detail state changed: {field}')
    old_items = {row['detail_id']: row for row in original['registry']}
    new_items = {row['detail_id']: row for row in incoming['registry']}
    if old_items.keys() != new_items.keys():
        raise ValueError('fixed detail IDs changed; use detail adoption for additions or identity changes')
    edit_fields = {'override', 'reason', 'source_ref', 'updated_at'}
    for detail_id, old in old_items.items():
        new = new_items[detail_id]
        changed = sorted(key for key in old.keys() | new.keys()
                         if key not in edit_fields and old.get(key) != new.get(key))
        if changed:
            raise ValueError(f'protected detail item changed: {detail_id}: {", ".join(changed)}')

    # Compare protected cells too: some visible columns are redundant with JSON
    # and are not represented by from_workbook's canonical result.
    def protected_rows(name, rows):
        result = []
        for index, row in enumerate(rows):
            row = list(row)
            if name == '詳細項目' and index:
                row += [None] * max(0, 16 - len(row))
                row = row[:6] + row[13:]
            elif name == '_detail_meta' and index and row[0] == 'revision':
                row[1] = None
            while row and row[-1] is None:
                row.pop()
            result.append(row)
        return result

    old_tables = workbook_rows(master_path, DETAIL_TABS)
    new_tables = workbook_rows(export_path, DETAIL_TABS)
    for name in DETAIL_TABS:
        if protected_rows(name, old_tables[name]) != protected_rows(name, new_tables[name]):
            raise ValueError(f'protected detail cells changed: {name}')

    def edit(row):
        return {field: copy.deepcopy(row.get(field, {} if field == 'override' else None))
                for field in sorted(edit_fields)}

    changes = []
    for detail_id in sorted(old_items):
        old, new = old_items[detail_id], new_items[detail_id]
        before, after = edit(old), edit(new)
        if before == after:
            continue
        fields = []
        for field in sorted(OVERRIDE_FIELDS):
            if old.get('override', {}).get(field) != new.get('override', {}).get(field):
                fields.append({'field': 'override.' + field,
                               'before': old.get('override', {}).get(field),
                               'after': new.get('override', {}).get(field)})
        for field in sorted(edit_fields - {'override'}):
            if before[field] != after[field]:
                fields.append({'field': field, 'before': before[field], 'after': after[field]})
        old_resolved, new_resolved = resolved_item(old), resolved_item(new)
        changes.append({
            'detail_id': detail_id, 'card_id': old['card_id'],
            'kind': old['source']['kind'], 'source_name': old['source']['name'],
            'before': before, 'after': after, 'field_changes': fields,
            'resolved_before': {field: old_resolved.get(field) for field in sorted(OVERRIDE_FIELDS)},
            'resolved_after': {field: new_resolved.get(field) for field in sorted(OVERRIDE_FIELDS)},
        })

    old_revision, new_revision = original['revision'], incoming['revision']
    if (type(new_revision) is not int or new_revision not in {old_revision, old_revision + 1}
            or (new_revision != old_revision and not changes)):
        raise ValueError('detail revision must stay unchanged or advance once with manual changes')
    old_edited = {key for key, row in old_items.items() if row.get('override')}
    new_edited = {key for key, row in new_items.items() if row.get('override')}
    return {
        'detail_changed': bool(changes),
        'detail_protected_unchanged': True,
        'detail_before_count': len(old_items),
        'detail_after_count': len(new_items),
        'detail_revision': {'before': old_revision, 'after': new_revision},
        'detail_changed_edit_ids': sorted(change['detail_id'] for change in changes
                                          if change['detail_id'] in old_edited & new_edited),
        'detail_removed_edit_ids': sorted(old_edited - new_edited),
        'detail_new_edit_ids': sorted(new_edited - old_edited),
        'detail_edit_changes': changes,
    }


def review(master_path, export_path):
    master_path, export_path = Path(master_path), Path(export_path)
    if master_path.resolve() == export_path.resolve():
        raise ValueError('export must be separate from the active master')
    protected_tabs = [name for name in TABS if name != '追加・手修正']
    old_tables = workbook_rows(master_path, protected_tabs)
    new_tables = workbook_rows(export_path, protected_tabs)
    for name in protected_tabs:
        if old_tables[name] != new_tables[name]:
            raise ValueError(f'protected tab changed: {name}')
    original, imported = load_master(master_path), load_master(export_path)
    for field in ['revision', 'coverage', 'source', 'evidence', 'runs', 'dictionaries', 'candidates']:
        if original[field] != imported[field]:
            raise ValueError(f'protected master state changed: {field}')
    old_registry = {r['card_id']: r for r in original['registry']}
    new_registry = {r['card_id']: r for r in imported['registry']}
    for card_id, record in old_registry.items():
        if new_registry.get(card_id) != record:
            raise ValueError(f'fixed ID or source aliases changed: {card_id}')
    added = sorted(set(new_registry) - set(old_registry))
    if any(new_registry[c]['aliases'] or new_registry[c]['family_id'] != c for c in added):
        raise ValueError('new manual card has an unexpected source alias or family')
    detail_review = review_details(master_path, export_path)
    before, after = resolve(original), resolve(imported)
    def identity(card):
        return (card['idol_id'], card['card_kind'], card['variant_kind'], card['card_title'])
    existing = {identity(c) for c in before}
    if any(identity(c) in existing for c in after if c['card_id'] in added):
        raise ValueError('possible manual duplicate; use the existing card ID')
    old_edits = {x['card_id']: x for x in original['overrides']}
    new_edits = {x['card_id']: x for x in imported['overrides']}
    return {
        'master_sha256': file_hash(master_path),
        'export_sha256': file_hash(export_path),
        'master_revision': original['revision'],
        'before_count': len(before),
        'after_count': len(after),
        'added_ids': added,
        'changed_edit_ids': sorted(c for c in new_edits.keys() & old_edits.keys() if new_edits[c] != old_edits[c]),
        'removed_edit_ids': sorted(old_edits.keys() - new_edits.keys()),
        'new_edit_ids': sorted(new_edits.keys() - old_edits.keys()),
        'edit_changes': [
            {'card_id': c, 'before': old_edits.get(c), 'after': new_edits.get(c)}
            for c in sorted(old_edits.keys() | new_edits.keys())
            if old_edits.get(c) != new_edits.get(c)
        ],
        'added_cards': [
            {k: c[k] for k in ['card_id', 'card_title', 'idol_id', 'idol_name', 'card_kind', 'variant_kind', 'wiki_url']}
            for c in after if c['card_id'] in added
        ],
        'protected_tabs_unchanged': True,
        **detail_review,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('mode', choices=['review', 'apply'])
    parser.add_argument('master')
    parser.add_argument('export')
    parser.add_argument('report')
    args = parser.parse_args()
    report_path = Path(args.report)
    if args.mode == 'review':
        if report_path.exists():
            raise FileExistsError('review report already exists; choose a new path')
        result = review(args.master, args.export)
        write(report_path, result)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return
    approved = read(report_path)
    current = review(args.master, args.export)
    if current != approved:
        raise ValueError('master or export changed since review; make a new review')
    if not current['detail_changed'] and not any(current[k] for k in ['added_ids', 'changed_edit_ids', 'removed_edit_ids', 'new_edit_ids']):
        print('No manual changes; active master was not rewritten.')
        return
    imported = load_master(args.export)
    save_master(imported, args.master, detail_source=args.export)
    if load_master(args.master) != imported:
        raise ValueError('applied master read-back mismatch')
    if current['detail_revision'] is not None:
        from src.detail_master import from_workbook
        if from_workbook(args.master) != from_workbook(args.export):
            raise ValueError('applied detail master read-back mismatch')
    print(json.dumps({'applied_at': now(), 'result': current}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
